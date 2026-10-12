"""完整知识导入与检索；模型空间不可变，异步任务完成后原子发布版本。"""

from __future__ import annotations

import asyncio
import io
from datetime import timedelta
from pathlib import Path
from typing import Any
from uuid import uuid4

from sqlalchemy import delete, select

from app.core.errors import AppError
from app.core.telemetry import retrieval
from app.db.models import now
from app.db.repository import attributes, owned, serialize
from app.modules.background.models import Outbox
from app.modules.knowledge.models import (
    Chunk,
    Document,
    ImportJob,
    KnowledgeBase,
    KnowledgeIndex,
)
from app.modules.storage.models import Artifact


class KnowledgeService:
    """与 HTTP 及 Agent 解耦的知识服务；解析器与重排器可由项目注册。"""

    def __init__(self, resources):
        self.resources, self.db = resources, resources.db
        self.parsers = {}
        self.reranker = None

    def register_parser(self, extension: str, parser):
        """注册真实格式解析器；返回含 content 与 page 的片段列表。"""
        self.parsers[extension.lower()] = parser

    def register_reranker(self, reranker):
        """接入项目真实重排器；接收 query 和带来源结果，返回重排结果。"""
        self.reranker = reranker

    def embedding_profile(self, name: str):
        """快照模型空间标识，不包含凭据；参数变化需要重建新空间。"""
        p = self.resources.models.profile(name, embedding=True)
        return {
            "name": name,
            "provider": p.provider,
            "model": p.model,
            "dimensions": p.dimensions,
            "endpoint": p.endpoint,
            "parameters": p.parameters,
        }

    async def create(self, principal, name: str, embedding: str):
        """创建空知识库与已声明模型空间；首次导入前无需调用收费模型。"""
        profile = self.embedding_profile(embedding)
        async with self.db.transaction(principal) as session:
            row = KnowledgeBase(**attributes(principal), name=name)
            session.add(row)
            await session.flush()
            index = KnowledgeIndex(
                **attributes(principal), knowledge_id=row.id, profile=profile, status="active"
            )
            session.add(index)
            await session.flush()
            row.active_index = index.id
            return serialize(row)

    async def submit(
        self,
        principal,
        knowledge_id: str,
        artifact_id: str,
        name: str,
        chunk_size=1200,
        overlap=150,
    ):
        """验证归属并记录导入及 outbox；重复内容和配置返回已有导入结果。"""
        if overlap < 0 or overlap >= chunk_size or chunk_size > 10000:
            raise AppError("chunk_configuration", "切分参数无效")
        async with self.db.transaction(principal) as session:
            kb = (
                await session.scalars(
                    select(KnowledgeBase)
                    .where(KnowledgeBase.id == knowledge_id, *owned(KnowledgeBase, principal))
                    .with_for_update()
                )
            ).one_or_none()
            artifact = (
                await session.scalars(
                    select(Artifact).where(Artifact.id == artifact_id, *owned(Artifact, principal))
                )
            ).one_or_none()
            if kb is None or artifact is None:
                raise AppError("resource_not_found", "知识库或附件不存在", 404)
            building = (
                await session.scalars(
                    select(KnowledgeIndex).where(
                        KnowledgeIndex.knowledge_id == kb.id, KnowledgeIndex.status == "building"
                    )
                )
            ).first()
            if building:
                raise AppError("knowledge_rebuilding", "索引重建期间不能修改文档", 409)
            config = {"chunk_size": chunk_size, "overlap": overlap, "sha256": artifact.sha256}
            jobs = (
                await session.scalars(
                    select(ImportJob).where(
                        ImportJob.knowledge_id == kb.id,
                        ImportJob.index_id == kb.active_index,
                        ImportJob.name == name,
                        ImportJob.status.in_(["queued", "running", "succeeded"]),
                        *owned(ImportJob, principal),
                    )
                )
            ).all()
            for existing in jobs:
                if existing.config == config:
                    return serialize(existing)
            job = ImportJob(
                **attributes(principal),
                knowledge_id=kb.id,
                index_id=kb.active_index,
                artifact_id=artifact.id,
                name=name,
                config=config,
            )
            session.add(job)
            await session.flush()
            session.add(Outbox(**attributes(principal), task="knowledge", resource_id=job.id))
            return serialize(job)

    async def import_action(self, principal, knowledge_id: str, job_id: str, action: str):
        """取消或重试导入任务；只有失败与已取消的任务允许重试，重试时重新写入 outbox。"""
        async with self.db.transaction(principal) as session:
            job = (
                await session.scalars(
                    select(ImportJob)
                    .where(
                        ImportJob.id == job_id,
                        ImportJob.knowledge_id == knowledge_id,
                        *owned(ImportJob, principal),
                    )
                    .with_for_update()
                )
            ).one_or_none()
            if job is None:
                raise AppError("resource_not_found", "导入任务不存在", 404)
            if action == "cancel":
                job.cancel_requested = True
                if job.status == "queued":
                    job.status = "cancelled"
            elif action == "retry" and job.status in {"failed", "cancelled"}:
                job.status, job.cancel_requested, job.error = "queued", False, None
                session.add(Outbox(**attributes(principal), task="knowledge", resource_id=job.id))
            else:
                raise AppError("import_conflict", "导入状态或操作不允许", 409)
            return serialize(job)

    async def artifact_in_use(self, session, artifact_id: str) -> bool:
        """附件仍被导入记录引用时返回 True；附件服务据此拒绝删除。"""
        reference = await session.scalars(
            select(ImportJob.id).where(ImportJob.artifact_id == artifact_id).limit(1)
        )
        return reference.first() is not None

    def parse(self, filename: str, data: bytes) -> list[dict[str, Any]]:
        """解析五类真实文档；扫描 PDF 没有文字时报告 OCR，禁止空内容成功。"""
        extension = Path(filename).suffix.lower()
        if extension in self.parsers:
            return self.parsers[extension](data)
        if extension in {".txt", ".md", ".markdown"}:
            return [{"content": data.decode("utf-8-sig"), "page": None}]
        if extension in {".html", ".htm"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(data, "html.parser")
            for tag in soup(["script", "style"]):
                tag.decompose()
            return [{"content": soup.get_text("\n", strip=True), "page": None}]
        if extension == ".docx":
            from docx import Document as WordDocument

            document = WordDocument(io.BytesIO(data))
            lines = [p.text for p in document.paragraphs]
            lines += [
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            return [{"content": "\n".join(lines), "page": None}]
        if extension == ".pdf":
            from pypdf import PdfReader

            document = PdfReader(io.BytesIO(data))
            pages = [
                {"content": page.extract_text() or "", "page": index + 1}
                for index, page in enumerate(document.pages)
            ]
            if not any(str(p["content"]).strip() for p in pages):
                raise AppError("ocr_required", "扫描 PDF 需要配置 OCR 解析器")
            return pages
        raise AppError("document_format", "文档格式不支持")

    def split(self, pages, size: int, overlap: int):
        """按页面保留定位并使用重叠字符窗口；空白片段不进入向量库。"""
        chunks = []
        for page in pages:
            content = page["content"]
            if not isinstance(content, str):
                raise AppError("document_parser", "解析器必须返回文本内容")
            for start in range(0, len(content), size - overlap):
                segment = content[start : start + size].strip()
                if segment:
                    chunks.append({"content": segment, "page": page["page"], "offset": start})
        if not chunks:
            raise AppError("document_empty", "文档没有可索引文本")
        return chunks

    async def ingest(self, job_id: str):
        """消费导入任务；重复投递跳过，网络调用在事务外，最终替换片段在短事务中完成。"""
        from app.core.security import Principal

        async with self.db.transaction(system=True) as session:
            job = (
                await session.scalars(
                    select(ImportJob).where(ImportJob.id == job_id).with_for_update()
                )
            ).one_or_none()
            if job is None or job.status != "queued":
                return
            if job.cancel_requested:
                job.status = "cancelled"
                return
            principal = Principal(job.tenant_id, job.subject_id, frozenset())
            job.status = "running"
            token = str(uuid4())
            job.lease_token = token
            job.lease_until = now() + timedelta(seconds=self.resources.settings.lease_seconds)
            index = await session.get(KnowledgeIndex, job.index_id)
            artifact_id, config, name = job.artifact_id, job.config, job.name
            profile, index_id, kb_id = index.profile, index.id, job.knowledge_id
        execution = asyncio.current_task()
        heartbeat = asyncio.create_task(self.renew(job_id, token, execution))
        try:
            if self.embedding_profile(profile["name"]) != profile:
                raise AppError("embedding_profile_changed", "Embedding 配置改变，需要先重建索引")
            artifact, data = await self.resources.artifacts.download(principal, artifact_id)
            pages = await asyncio.to_thread(self.parse, artifact["filename"], data)
            chunks = self.split(pages, config["chunk_size"], config["overlap"])
            vectors = []
            for start in range(0, len(chunks), 32):
                async with self.db.transaction(principal) as session:
                    job = await session.get(ImportJob, job_id)
                    if job.lease_token != token:
                        raise AppError("lease_lost", "导入租约已失效", 409)
                    if job.cancel_requested:
                        job.status = "cancelled"
                        return
                    job.progress = int(100 * start / len(chunks))
                vectors.extend(
                    await self.resources.models.embed(
                        profile["name"], [c["content"] for c in chunks[start : start + 32]]
                    )
                )
            async with self.db.transaction(principal) as session:
                job = (
                    await session.scalars(
                        select(ImportJob).where(ImportJob.id == job_id).with_for_update()
                    )
                ).one()
                kb = (
                    await session.scalars(
                        select(KnowledgeBase).where(KnowledgeBase.id == kb_id).with_for_update()
                    )
                ).one()
                if job.lease_token != token:
                    raise AppError("lease_lost", "导入租约已失效", 409)
                if job.cancel_requested:
                    job.status = "cancelled"
                    return
                document = (
                    await session.scalars(
                        select(Document).where(
                            Document.knowledge_id == kb_id, Document.name == name
                        )
                    )
                ).one_or_none()
                if document is None:
                    document = Document(**attributes(principal), knowledge_id=kb_id, name=name)
                    session.add(document)
                    await session.flush()
                else:
                    document.version += 1
                await session.execute(
                    delete(Chunk).where(
                        Chunk.document_id == document.id, Chunk.index_id == index_id
                    )
                )
                for position, (chunk, vector) in enumerate(zip(chunks, vectors, strict=True)):
                    session.add(
                        Chunk(
                            **attributes(principal),
                            document_id=document.id,
                            index_id=index_id,
                            version=document.version,
                            position=position,
                            page=chunk["page"],
                            content=chunk["content"],
                            embedding=vector,
                            metadata_={
                                "offset": chunk["offset"],
                                "name": name,
                                "artifact_id": artifact_id,
                                "sha256": artifact["sha256"],
                                "parser": Path(artifact["filename"]).suffix,
                            },
                        )
                    )
                job.status, job.progress = "succeeded", 100
                job.lease_token, job.lease_until = None, None
                # 重建的全部导入都完成时才发布；一个失败就保持旧空间可用。
                if index_id != kb.active_index:
                    pending = (
                        await session.scalars(
                            select(ImportJob).where(
                                ImportJob.index_id == index_id,
                                ImportJob.id != job_id,
                                ImportJob.status != "succeeded",
                            )
                        )
                    ).first()
                    if pending is None:
                        old = await session.get(KnowledgeIndex, kb.active_index)
                        old.status = "retired"
                        index = await session.get(KnowledgeIndex, index_id)
                        index.status, kb.active_index = "active", index_id
        except Exception as error:
            async with self.db.transaction(principal) as session:
                job = await session.get(ImportJob, job_id)
                if job and job.lease_token == token:
                    job.status = "failed"
                    job.lease_token, job.lease_until = None, None
                    job.error = {
                        "code": error.code if isinstance(error, AppError) else "ingestion_failed",
                        "message": error.message if isinstance(error, AppError) else "文档导入失败",
                    }
            raise
        finally:
            heartbeat.cancel()
            await asyncio.gather(heartbeat, return_exceptions=True)

    async def renew(self, job_id, token, execution):
        """续租与取消均通过数据库判断；过期实例停止消费，不能覆盖新实例结果。"""
        while execution and not execution.done():
            await asyncio.sleep(self.resources.settings.lease_seconds / 3)
            async with self.db.transaction(system=True) as session:
                job = await session.get(ImportJob, job_id)
                if not job or job.lease_token != token or job.cancel_requested:
                    if job and job.lease_token == token and job.cancel_requested:
                        job.status = "cancelled"
                        job.lease_token, job.lease_until = None, None
                    execution.cancel()
                    return
                job.lease_until = now() + timedelta(seconds=self.resources.settings.lease_seconds)

    async def recover(self):
        """崩溃导入失效后重新排队；片段原子发布，重复收费请求只发生在未提交批次。"""
        async with self.db.transaction(system=True) as session:
            jobs = (
                await session.scalars(
                    select(ImportJob)
                    .where(ImportJob.status == "running", ImportJob.lease_until < now())
                    .with_for_update(skip_locked=True)
                )
            ).all()
            for job in jobs:
                job.status, job.lease_token, job.lease_until = "queued", None, None
                session.add(
                    Outbox(
                        tenant_id=job.tenant_id,
                        subject_id=job.subject_id,
                        task="knowledge",
                        resource_id=job.id,
                    )
                )
        return len(jobs)

    async def rebuild(self, principal, knowledge_id: str, embedding: str):
        """创建新空间并重投现有原文件；旧空间在所有文档成功前持续可查询。"""
        profile = self.embedding_profile(embedding)
        async with self.db.transaction(principal) as session:
            kb = (
                await session.scalars(
                    select(KnowledgeBase)
                    .where(KnowledgeBase.id == knowledge_id, *owned(KnowledgeBase, principal))
                    .with_for_update()
                )
            ).one_or_none()
            if kb is None:
                raise AppError("resource_not_found", "知识库不存在", 404)
            busy = (
                await session.scalars(
                    select(ImportJob).where(
                        ImportJob.knowledge_id == kb.id, ImportJob.status.in_(["queued", "running"])
                    )
                )
            ).first()
            building = (
                await session.scalars(
                    select(KnowledgeIndex).where(
                        KnowledgeIndex.knowledge_id == kb.id, KnowledgeIndex.status == "building"
                    )
                )
            ).first()
            if busy or building:
                raise AppError("knowledge_busy", "知识库有导入或重建任务", 409)
            index = KnowledgeIndex(**attributes(principal), knowledge_id=kb.id, profile=profile)
            session.add(index)
            await session.flush()
            docs = (
                await session.scalars(select(Document).where(Document.knowledge_id == kb.id))
            ).all()
            for document in docs:
                chunk = (
                    await session.scalars(
                        select(Chunk)
                        .where(Chunk.document_id == document.id, Chunk.index_id == kb.active_index)
                        .limit(1)
                    )
                ).first()
                if chunk is None:
                    continue
                job = ImportJob(
                    **attributes(principal),
                    knowledge_id=kb.id,
                    index_id=index.id,
                    artifact_id=chunk.metadata_["artifact_id"],
                    name=document.name,
                    config={
                        "chunk_size": 1200,
                        "overlap": 150,
                        "sha256": chunk.metadata_["sha256"],
                    },
                )
                session.add(job)
                await session.flush()
                session.add(Outbox(**attributes(principal), task="knowledge", resource_id=job.id))
            if not docs:
                index.status, kb.active_index = "active", index.id
            return serialize(index)

    async def search(
        self, principal, knowledge_id: str, query: str, top_k=5, metadata=None, rerank=False
    ):
        """只查询当前主体的活动空间；返回内容、距离与文档版本和定位出处。"""
        with retrieval.time():
            async with self.db.transaction(principal) as session:
                kb = (
                    await session.scalars(
                        select(KnowledgeBase).where(
                            KnowledgeBase.id == knowledge_id, *owned(KnowledgeBase, principal)
                        )
                    )
                ).one_or_none()
                if kb is None:
                    raise AppError("resource_not_found", "知识库不存在", 404)
                index = await session.get(KnowledgeIndex, kb.active_index)
                profile, index_id = index.profile, index.id
            if self.embedding_profile(profile["name"]) != profile:
                raise AppError("embedding_profile_changed", "Embedding 配置改变，需要重建索引")
            vector = (await self.resources.models.embed(profile["name"], [query]))[0]
            from pgvector.sqlalchemy import HALFVEC, Vector
            from sqlalchemy import cast

            dimensions = profile["dimensions"]
            expression = (
                cast(Chunk.embedding, Vector(dimensions))
                if dimensions <= 2000
                else cast(Chunk.embedding, HALFVEC(dimensions))
                if dimensions <= 4000
                else Chunk.embedding
            )
            distance = expression.cosine_distance(vector).label("distance")
            async with self.db.transaction(principal) as session:
                from sqlalchemy import func, text

                await session.execute(text("SET LOCAL hnsw.iterative_scan = strict_order"))
                stmt = select(Chunk, distance).where(
                    Chunk.index_id == index_id,
                    func.vector_dims(Chunk.embedding) == dimensions,
                    *owned(Chunk, principal),
                )
                if metadata:
                    stmt = stmt.where(Chunk.metadata_.contains(metadata))
                rows = (
                    await session.execute(
                        stmt.order_by(distance).limit(top_k * 3 if rerank else top_k)
                    )
                ).all()
                results = [
                    {
                        "content": chunk.content,
                        "distance": float(score),
                        "document_id": chunk.document_id,
                        "version": chunk.version,
                        "position": chunk.position,
                        "page": chunk.page,
                        "metadata": chunk.metadata_,
                    }
                    for chunk, score in rows
                ]
            if rerank:
                if self.reranker is None:
                    raise AppError("reranker_not_configured", "重排器未注册", 503)
                results = await self.reranker(query, results)
            return results[:top_k]

    async def delete_document(self, principal, document_id: str):
        """先锁知识库再删除文档；运行中的导入和索引重建不允许被删除竞态打断。"""
        async with self.db.transaction(principal) as session:
            document = (
                await session.scalars(
                    select(Document).where(Document.id == document_id, *owned(Document, principal))
                )
            ).one_or_none()
            if document is None:
                raise AppError("resource_not_found", "文档不存在", 404)
            await session.execute(
                select(KnowledgeBase)
                .where(KnowledgeBase.id == document.knowledge_id)
                .with_for_update()
            )
            active = (
                await session.scalars(
                    select(ImportJob).where(
                        ImportJob.knowledge_id == document.knowledge_id,
                        ImportJob.name == document.name,
                        ImportJob.status.in_(["queued", "running"]),
                    )
                )
            ).first()
            building = (
                await session.scalars(
                    select(KnowledgeIndex).where(
                        KnowledgeIndex.knowledge_id == document.knowledge_id,
                        KnowledgeIndex.status == "building",
                    )
                )
            ).first()
            if active or building:
                raise AppError("knowledge_busy", "导入或重建期间不能删除文档", 409)
            await session.execute(
                delete(ImportJob).where(
                    ImportJob.knowledge_id == document.knowledge_id, ImportJob.name == document.name
                )
            )
            await session.delete(document)

    async def delete_knowledge(self, principal, knowledge_id: str):
        """仅空闲知识库可删除；原文件附件独立保留，可由附件接口清理。"""
        async with self.db.transaction(principal) as session:
            kb = (
                await session.scalars(
                    select(KnowledgeBase)
                    .where(KnowledgeBase.id == knowledge_id, *owned(KnowledgeBase, principal))
                    .with_for_update()
                )
            ).one_or_none()
            if kb is None:
                raise AppError("resource_not_found", "知识库不存在", 404)
            active = (
                await session.scalars(
                    select(ImportJob).where(
                        ImportJob.knowledge_id == kb.id, ImportJob.status.in_(["queued", "running"])
                    )
                )
            ).first()
            if active:
                raise AppError("knowledge_busy", "知识库存在活动导入", 409)
            await session.delete(kb)

    def langchain_retriever(self, principal, knowledge_id: str, top_k=5):
        """构造原生 BaseRetriever；LangGraph 可直接注入服务或此对象。"""
        from langchain_core.documents import Document as LangDocument
        from langchain_core.retrievers import BaseRetriever

        service = self

        class Retriever(BaseRetriever):
            """受身份约束的知识检索器，只开放异步网络路径。"""

            def _get_relevant_documents(self, query, *, run_manager):
                raise AppError("async_required", "检索必须使用异步接口")

            async def _aget_relevant_documents(self, query, *, run_manager):
                rows = await service.search(principal, knowledge_id, query, top_k)
                return [
                    LangDocument(
                        page_content=row["content"],
                        metadata={k: v for k, v in row.items() if k != "content"},
                    )
                    for row in rows
                ]

        return Retriever()
